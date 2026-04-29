from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy

RED = RGBColor(0xCC, 0x00, 0x00)
DARK_RED = RGBColor(0xA0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color=None, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width or 1)
    return shape


def add_textbox(slide, left, top, width, height, text="", font_size=18, bold=False,
                color=BLACK, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_paragraph(text_frame, text, font_size=14, bold=False, color=BLACK, alignment=PP_ALIGN.LEFT,
                  space_before=Pt(2), space_after=Pt(2), level=0, font_name="Calibri"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = space_before
    p.space_after = space_after
    p.level = level
    return p


def add_bullet(text_frame, text, font_size=14, bold=False, color=BLACK, level=0, font_name="Calibri"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.level = level
    p.space_before = Pt(3)
    p.space_after = Pt(3)
    return p


def add_title_bar(slide, title_text):
    add_shape(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.9), fill_color=RED)
    add_textbox(slide, Inches(0.6), Inches(0.12), Inches(12), Inches(0.7),
                title_text, font_size=28, bold=True, color=WHITE)


def add_table(slide, rows, cols, data, left, top, width, height, col_widths=None, font_size=11):
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c] if r < len(data) and c < len(data[r]) else ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(font_size)
                paragraph.font.name = "Calibri"
                if r == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = WHITE
                else:
                    paragraph.font.color.rgb = BLACK

            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RED
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE

    return table_shape


# =========================================================
# SLIDE 1: Title
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), prs.slide_width, Inches(3.0), fill_color=RED)
add_textbox(slide, Inches(1), Inches(0.8), Inches(11), Inches(1.2),
            "Red Hat Architecture", font_size=44, bold=True, color=WHITE)
add_textbox(slide, Inches(1), Inches(2.0), Inches(11), Inches(0.8),
            "30-Day Scope \u2014 Outcomes", font_size=30, bold=False, color=WHITE)
add_textbox(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.6),
            "RACI Matrices  \u2022  Role Charters  \u2022  Skills Matrix  \u2022  EA Remit", font_size=20, color=GRAY)
add_textbox(slide, Inches(1), Inches(5.0), Inches(11), Inches(0.5),
            "Scaled Agile Framework  \u2022  2026", font_size=16, color=GRAY)

# =========================================================
# SLIDE 2: Agenda
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Agenda")

agenda_items = [
    "Architecture RACI Matrix",
    "Role Definitions & Interaction Mappings",
    "Key Differentiators Across Roles",
    "EA Function Remit (Services to Red Hat)",
    "Data Architecture Discipline",
    "Platform Architect Role",
    "Architecture Skills Matrix",
    "Needed Inputs for Enterprise Architecture",
]
txBox = add_textbox(slide, Inches(1.2), Inches(1.3), Inches(10), Inches(5.5), "", font_size=20)
tf = txBox.text_frame
tf.word_wrap = True
for i, item in enumerate(agenda_items):
    add_paragraph(tf, f"{i+1}.   {item}", font_size=22, color=BLACK, space_before=Pt(10), space_after=Pt(10))

# =========================================================
# SLIDE 3: RACI — Strategy & Governance
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "RACI Matrix \u2014 Strategy & Governance")

data = [
    ["Task / Deliverable", "EA", "DA", "SA", "BA", "PdM"],
    ["Define 1\u20133+ Year IT Strategy & North Star", "A / R", "C", "C", "C", "C"],
    ["Create Capability Briefs & Value Stream Maps", "C", "C", "I", "A / R", "C"],
    ["Align Architecture with OKRs & Org Strategy", "A / R", "I", "I", "C", "C"],
    ["Define AI-Ready Architecture Principles", "C", "A / R", "C", "I", "I"],
    ["Data Quality, Governance & Explainability", "R", "A / R", "C", "I", "I"],
    ["Set Enterprise Technology Radar & Standards", "A / R", "C", "C", "I", "I"],
]
add_table(slide, len(data), 6, data,
          Inches(0.5), Inches(1.2), Inches(12.3), Inches(4.2),
          col_widths=[Inches(4.8), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5)],
          font_size=13)

add_textbox(slide, Inches(0.6), Inches(5.8), Inches(12), Inches(0.6),
            "R = Responsible    A = Accountable    C = Consulted    I = Informed",
            font_size=13, color=GRAY)

# =========================================================
# SLIDE 4: RACI — Solution Design & Lifecycle
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "RACI Matrix \u2014 Solution Design & Lifecycle")

data = [
    ["Task / Deliverable", "EA", "DA", "SA", "BA", "PdM"],
    ["High-Level Technical Vision & Architecture", "C", "C", "A / R", "C", "C"],
    ["Architectural Decision Records (ADR)", "C / I", "I", "A / R", "I", "I"],
    ["Conduct Proof of Concepts (POCs)", "C", "C", "A / R", "I", "I"],
    ["Solution Architecture Document (SAD)", "C", "I", "A / R", "I", "I"],
    ["Product Roadmap Prioritization", "C", "I", "C", "C", "A / R"],
    ["Implementation Handover Ceremony", "I", "I", "R", "I", "A"],
    ["Continuous Improvement Framework", "A / R", "C", "C", "C", "C"],
]
add_table(slide, len(data), 6, data,
          Inches(0.5), Inches(1.2), Inches(12.3), Inches(4.8),
          col_widths=[Inches(4.8), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5)],
          font_size=13)

add_textbox(slide, Inches(0.6), Inches(6.3), Inches(12), Inches(0.6),
            "R = Responsible    A = Accountable    C = Consulted    I = Informed",
            font_size=13, color=GRAY)

# =========================================================
# SLIDE 5: Role Definitions — EA & SA
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Role Definitions")

box1 = add_shape(slide, Inches(0.5), Inches(1.2), Inches(5.9), Inches(5.6), line_color=RED, line_width=2)
txBox = add_textbox(slide, Inches(0.8), Inches(1.3), Inches(5.4), Inches(5.4), "", font_size=14)
tf = txBox.text_frame
tf.word_wrap = True
add_paragraph(tf, "Enterprise & Data Architect", font_size=22, bold=True, color=RED)
add_paragraph(tf, '"The Foundation Visionary"', font_size=14, color=GRAY, space_after=Pt(12))
add_bullet(tf, "Operates at Strategy Office level", font_size=14)
add_bullet(tf, "Ensures technical roadmaps enable business outcomes", font_size=14)
add_bullet(tf, "Output: Technology Radar, Data Architecture standards, target-state capability roadmaps", font_size=14)
add_bullet(tf, "Advisory authority on Buy vs. Build & vendor alignment", font_size=14)
add_bullet(tf, "AI Focus: Data quality & explainability for Agentic AI", font_size=14)

box2 = add_shape(slide, Inches(6.9), Inches(1.2), Inches(5.9), Inches(5.6), line_color=BLUE, line_width=2)
txBox = add_textbox(slide, Inches(7.2), Inches(1.3), Inches(5.4), Inches(5.4), "", font_size=14)
tf = txBox.text_frame
tf.word_wrap = True
add_paragraph(tf, "Solution Architect", font_size=22, bold=True, color=BLUE)
add_paragraph(tf, '"The Technical Visionary"', font_size=14, color=GRAY, space_after=Pt(12))
add_bullet(tf, "Operates at Portfolio & Project level", font_size=14)
add_bullet(tf, 'Bridges BA\'s "what/why" \u2192 technical "how"', font_size=14)
add_bullet(tf, "Output: SADs, ADRs, successful POCs", font_size=14)
add_bullet(tf, "Project-specific tech choices within EA/DA guardrails", font_size=14)
add_bullet(tf, "Leads Implementation Handover Ceremony", font_size=14)

# =========================================================
# SLIDE 6: Data Architect Role
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Data Architect Role")

box = add_shape(slide, Inches(1.5), Inches(1.3), Inches(10.3), Inches(5.5), line_color=GREEN, line_width=2)
txBox = add_textbox(slide, Inches(1.9), Inches(1.5), Inches(9.5), Inches(5.2), "", font_size=14)
tf = txBox.text_frame
tf.word_wrap = True
add_paragraph(tf, "Data Architect (DA)", font_size=24, bold=True, color=GREEN)
add_paragraph(tf, '"The Information Steward"', font_size=15, color=GRAY, space_after=Pt(14))
add_bullet(tf, "Strategic Scope: Enterprise & portfolio levels \u2014 defines how data is modeled, integrated, governed, and made usable across business capabilities", font_size=15, level=0)
add_bullet(tf, "Primary Output: Conceptual & logical data models, data standards, integration patterns, target-state data flows", font_size=15, level=0)
add_bullet(tf, "Decision Rights: Advisory authority on data design, master data approach, integration patterns & governance requirements", font_size=15, level=0)
add_bullet(tf, "AI Focus: Trusted data foundation for AI \u2014 data quality, lineage, interoperability, and fitness for AI-driven workflows", font_size=15, level=0)

# =========================================================
# SLIDE 7: Key Workflow Interlocks
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Key Workflow Interlocks")

interlocks = [
    ("Triggering Engagement", "When Corporate Strategy incubates a business outcome, they notify both EA and Business Architects to determine whether the impact rises to EA or SA engagement."),
    ("Cadence Alignment", "All architecture teams operate on a 10-week increment cadence to ensure synchronization with PI Planning."),
    ("Strategic Integration", 'EA is being moved from a "downstream service" to an upstream partner, gaining a "seat at the table" during the investment lifecycle to prevent misaligned investments.'),
]
y = Inches(1.4)
for title, desc in interlocks:
    add_shape(slide, Inches(0.8), y, Inches(0.12), Inches(1.3), fill_color=RED)
    box = add_shape(slide, Inches(0.92), y, Inches(11.5), Inches(1.3), fill_color=RGBColor(0xFE, 0xF2, 0xF2))
    txBox = add_textbox(slide, Inches(1.2), y + Inches(0.1), Inches(11), Inches(1.1), "", font_size=14)
    tf = txBox.text_frame
    tf.word_wrap = True
    add_paragraph(tf, title, font_size=17, bold=True, color=RED)
    add_paragraph(tf, desc, font_size=14, color=BLACK, space_before=Pt(4))
    y += Inches(1.8)

# =========================================================
# SLIDE 8: Key Differentiators (1 of 2)
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Key Differentiators (1 of 2)")

data = [
    ["Feature", "Enterprise Architect", "Solution Architect", "Data Architect", "Business Architect"],
    ["Scope", "Entire org & portfolio;\nenterprise guardrails & North Star", "Specific solution/project;\nend-to-end technical solution", "Enterprise data domains;\ncross-platform data ecosystem", "Business capabilities\n& value streams (CBP)"],
    ["Time Horizon", "1\u20133+ years (strategic)", "3\u201312 months (tactical)", "1\u20133 years\n(strategic to transitional)", "1\u20133+ years\n(strategic to translational)"],
    ["North Star", "AI-ready architecture;\nreduce sprawl & enable reuse", "Feasible solution design\nmeeting NFR targets", "Trusted, governed,\nexplainable data at scale", "Coherent capability map\nrealizing OKRs"],
    ["Success Metric", "Portfolio ROI, reduced\nredundancy, strategic alignment", "Delivery success, system\nperformance, technical fit", "Data quality, consistency,\nreuse, trusted AI", "Capability realization,\nvalue stream performance"],
]
add_table(slide, len(data), 5, data,
          Inches(0.3), Inches(1.15), Inches(12.7), Inches(5.5),
          col_widths=[Inches(2.0), Inches(2.675), Inches(2.675), Inches(2.675), Inches(2.675)],
          font_size=11)

# =========================================================
# SLIDE 9: Key Differentiators (2 of 2)
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Key Differentiators (2 of 2)")

data = [
    ["Feature", "Enterprise Architect", "Solution Architect", "Data Architect", "Business Architect"],
    ["Risk Focus", "High: wrong platforms,\nmisaligned bets", "Medium: wrong patterns/\ntools, delivery risk", "High: inconsistent entities,\npoor lineage", "Med-high: misaligned\ncapabilities, CBP gaps"],
    ["Typical\nMeetings", "Strategy reviews, portfolio\ncouncils, ARB, radar sessions", "PI Planning, design reviews,\nspike/POC reviews, handover", "Governance councils, domain\nreviews, AI/data forums", "Capability workshops, VSM,\nBA councils, OKR sessions"],
    ["Key\nOutputs", "Tech Radar, EA roadmaps,\nAI-ready principles, ref archs", "SADs, ADRs, POC results,\nrunway updates", "Data models, canonical entities,\ndata standards, patterns", "Capability briefs, VSMs,\nRACIs, BDRs, CBP views"],
    ["SAFe\nAnchor", "Strategy Office, LPM,\nPlatform/Domain architects", "Bridges BA \u2192 technical;\nworks with Product, DA", "Interlocks EA, SA, BA,\ndata governance", "Engages Strategy & EA;\nshapes CBP & value streams"],
]
add_table(slide, len(data), 5, data,
          Inches(0.3), Inches(1.15), Inches(12.7), Inches(5.5),
          col_widths=[Inches(2.0), Inches(2.675), Inches(2.675), Inches(2.675), Inches(2.675)],
          font_size=11)

# =========================================================
# SLIDE 10: EA Remit — Strategic Partnership & Governance
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "EA Remit: Strategic Partnership & Governance")

items = [
    ("Strategy Office Representation", "Formal seat at the Strategy Office ensuring technical roadmaps enable business outcomes"),
    ("Investment Vetting", "Technology and portfolio investments vetted against business goals before capital commitment"),
    ("Decision Rights Management", 'Advisory authority on "Buy vs. Build" and vendor roadmap alignment'),
    ("AI Policy & Ethics", "Governing AI policy, security controls, and ethical frameworks for frictionless AI production"),
]
txBox = add_textbox(slide, Inches(1), Inches(1.4), Inches(11.3), Inches(5.5), "", font_size=16)
tf = txBox.text_frame
tf.word_wrap = True
for title, desc in items:
    add_bullet(tf, f"{title} \u2014 {desc}", font_size=17)

# =========================================================
# SLIDE 11: EA Remit — Standards & Foundation Building
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "EA Remit: Standards & Foundation Building")

items = [
    ("Data Architecture Discipline", "Standalone service for data standards, quality, stewardship, and governance"),
    ("AI-Ready Architecture Principles", "Explainability and interoperability for scalable AI & agentic workflows"),
    ("Technology Radar Management", "Red Hat Technical Radar by function/value stream to guide technology selection"),
    ("Reusable Patterns", "Reference architectures for common GTM scenarios (e.g., consumption pricing, AI-enabled selling)"),
]
txBox = add_textbox(slide, Inches(1), Inches(1.4), Inches(11.3), Inches(5.5), "", font_size=16)
tf = txBox.text_frame
tf.word_wrap = True
for title, desc in items:
    add_bullet(tf, f"{title} \u2014 {desc}", font_size=17)

# =========================================================
# SLIDE 12: EA Remit — Capability & Roadmap Design
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "EA Remit: Capability & Roadmap Design")

items = [
    ("Proactive Capability Design", "Shifting from reactive project delivery to proactive design anticipating future needs"),
    ("Strategic Roadmapping", "Enterprise-level roadmaps synthesizing solution roadmaps & identifying non-obvious dependencies"),
    ("Target-State Architecture", "\"North Star\" vision for RH's technology landscape over a 3\u20135 year horizon"),
    ("Interoperability Oversight", "Seamless integration and data flow between disparate systems and platforms"),
]
txBox = add_textbox(slide, Inches(1), Inches(1.4), Inches(11.3), Inches(5.5), "", font_size=16)
tf = txBox.text_frame
tf.word_wrap = True
for title, desc in items:
    add_bullet(tf, f"{title} \u2014 {desc}", font_size=17)

# =========================================================
# SLIDE 13: EA Remit — Portfolio & Operational Excellence
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "EA Remit: Portfolio & Operational Excellence")

items = [
    ("Architecture Review Board (ARB)", "Formal approval process for architectural standards & cross-portfolio alignment"),
    ("Portfolio Impact Assessment", "Architecture impact assessments for new requests to determine validity & fit"),
    ("Vendor Roadmap Evaluation", "Vendor capability roadmaps analyzed against internal gaps"),
    ("Continuous Improvement Framework", "Maintenance, monitoring, and feedback loops for sustained excellence"),
]
txBox = add_textbox(slide, Inches(1), Inches(1.4), Inches(11.3), Inches(5.5), "", font_size=16)
tf = txBox.text_frame
tf.word_wrap = True
for title, desc in items:
    add_bullet(tf, f"{title} \u2014 {desc}", font_size=17)

# =========================================================
# SLIDE 14: Data Architecture — What It Is vs. Isn't
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Data Architecture: What It Is vs. What It Isn\u2019t")

data = [
    ["Feature", "What It IS", "What It IS NOT"],
    ["Primary Goal", "Ensuring data quality, governance &\nexplainability for scalable AI", "Enforcing rigid barriers that\nstifle AI experimentation"],
    ["Focus", "Managing data standards,\nstewardship & data products", "General \"Enterprise Architecture\"\nor downstream delivery"],
    ["Approach", "Proactively defining data architecture\nprinciples & AI-ready standards", "A reactive \"firefighter\" called\nin at the last minute"],
    ["Output", "DA documentation, Data Product\nEpics & BDRs", "A mere inventory of systems\nor an implementation roadmap"],
]
tbl = add_table(slide, len(data), 3, data,
          Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.0),
          col_widths=[Inches(2.5), Inches(4.6), Inches(4.6)],
          font_size=13)

for r in range(1, len(data)):
    cell = tbl.table.cell(r, 2)
    for p in cell.text_frame.paragraphs:
        p.font.color.rgb = RGBColor(0x99, 0x33, 0x33)

# =========================================================
# SLIDE 15: Data & AI Architect — Role & Responsibilities
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Data & AI Architect: Role & Responsibilities")

txBox = add_textbox(slide, Inches(0.6), Inches(1.2), Inches(5.8), Inches(5.8), "", font_size=14)
tf = txBox.text_frame
tf.word_wrap = True
add_paragraph(tf, "Role Definition", font_size=20, bold=True, color=GREEN)
add_bullet(tf, "Purpose: Create high-level technical vision & data architecture achieving business goals", font_size=14)
add_bullet(tf, 'Engagement: Triggered by BA requests \u2192 creates Epics in Jira with label "Data and AI Architecture"', font_size=14)
add_bullet(tf, "Interlocks: SAs, Product Owners, TOC", font_size=14)
add_paragraph(tf, "", font_size=8)
add_paragraph(tf, "Required Skills", font_size=20, bold=True, color=GREEN)
add_bullet(tf, "Domain Expertise", font_size=14)
add_bullet(tf, "Data Governance", font_size=14)
add_bullet(tf, "Technical Proficiency", font_size=14)
add_bullet(tf, "Strategic Alignment", font_size=14)
add_bullet(tf, "AI Specialization (Agentic AI, MCP)", font_size=14)

txBox2 = add_textbox(slide, Inches(6.8), Inches(1.2), Inches(6), Inches(5.8), "", font_size=14)
tf2 = txBox2.text_frame
tf2.word_wrap = True
add_paragraph(tf2, "Core Responsibilities", font_size=20, bold=True, color=GREEN)
add_bullet(tf2, "Standardization \u2014 Enterprise data standards for consistency", font_size=14)
add_bullet(tf2, "AI Enablement \u2014 Platform-specific AI plans; AI embedded into workflows", font_size=14)
add_bullet(tf2, "Impact Assessment \u2014 CDE & EDS considerations for data products", font_size=14)
add_bullet(tf2, "Stewardship & Quality \u2014 Data governance & explainability for trusted AI", font_size=14)
add_bullet(tf2, "Lifecycle Mgmt \u2014 Tech evaluation & decisions for data solutions", font_size=14)

# =========================================================
# SLIDE 16: Platform Architect
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Platform Architect")

add_shape(slide, Inches(0.6), Inches(1.2), Inches(12.1), Inches(3.0), line_color=PURPLE, line_width=2)
txBox = add_textbox(slide, Inches(1.0), Inches(1.3), Inches(11.3), Inches(2.8), "", font_size=14)
tf = txBox.text_frame
tf.word_wrap = True
add_paragraph(tf, "Platform Architect (Agile / DevOps context)", font_size=20, bold=True, color=PURPLE)
add_bullet(tf, "Operates at a platform or shared-services level", font_size=15)
add_bullet(tf, "Designs & evolves technical platforms (cloud, K8s, CI/CD, shared APIs, data platforms)", font_size=15)
add_bullet(tf, "Focuses on scalability, reliability, security, multi-tenancy, cost optimization, developer experience", font_size=15)
add_bullet(tf, "Aligns platform capabilities with EA guardrails and solution/product team needs", font_size=15)
add_bullet(tf, "Does not own enterprise-wide business-technology strategy or full end-to-end solutions", font_size=15)

data = [
    ["Aspect", "Platform Architect"],
    ["Primary Level", "Shared platform / infrastructure / tooling"],
    ["Time Horizon", "Medium\u2013long term (platform evolution & lifecycle)"],
    ["Key Outputs", "Platform blueprints, SLOs, platform APIs, cloud/infra patterns"],
    ["Partners", "SRE/DevOps, infra teams, EAs, SAs, security"],
    ["SAFe Anchor", "Part of platform / system architecture (not a named SAFe role)"],
]
add_table(slide, len(data), 2, data,
          Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.6),
          col_widths=[Inches(3.0), Inches(8.7)],
          font_size=13)

# =========================================================
# SLIDE 17: Skills Comparison Matrix
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Skills Comparison Matrix")

data = [
    ["Skill Category", "EA", "SA", "DA", "BA"],
    ["Strategic Planning", "Expert", "Intermediate", "Advanced", "Expert"],
    ["Technical Depth", "Intermediate", "Expert", "Expert", "Basic"],
    ["Business Acumen", "Advanced", "Intermediate", "Intermediate", "Expert"],
    ["Governance", "Expert", "Intermediate", "Expert", "Intermediate"],
    ["Architecture Frameworks", "Expert", "Advanced", "Advanced", "Expert"],
    ["AI/ML Knowledge", "Advanced", "Intermediate", "Expert", "Basic"],
    ["Agile/SAFe", "Intermediate", "Expert", "Intermediate", "Advanced"],
    ["Stakeholder Mgmt", "Expert", "Advanced", "Advanced", "Expert"],
    ["Documentation", "Advanced", "Expert", "Expert", "Expert"],
    ["Integration Design", "Intermediate", "Expert", "Expert", "Intermediate"],
]

tbl = add_table(slide, len(data), 5, data,
          Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.2),
          col_widths=[Inches(3.5), Inches(2.2), Inches(2.2), Inches(2.2), Inches(2.2)],
          font_size=13)

prof_colors = {
    "Expert": RED,
    "Advanced": BLUE,
    "Intermediate": AMBER,
    "Basic": GRAY,
}
for r in range(1, len(data)):
    for c in range(1, 5):
        cell = tbl.table.cell(r, c)
        val = data[r][c]
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = prof_colors.get(val, BLACK)
            p.font.bold = True
            p.alignment = PP_ALIGN.CENTER

add_textbox(slide, Inches(0.6), Inches(6.5), Inches(12), Inches(0.5),
            "Expert = Define standards & mentor  |  Advanced = Independently execute  |  Intermediate = Contribute with guidance  |  Basic = Foundational",
            font_size=11, color=GRAY)

# =========================================================
# SLIDE 18: Skills Development Priorities
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Skills Development Priorities")

roles = [
    ("Enterprise Architects", RED, Inches(0.5), [
        "AI governance & explainability",
        "Portfolio financial modeling & ROI",
        "Multi-cloud architecture patterns",
        "Change leadership",
    ]),
    ("Solution Architects", BLUE, Inches(3.7), [
        "AI/ML implementation patterns",
        "Cloud-native & serverless expertise",
        "SAFe architecture runway & CD",
        "API strategy & platform architecture",
    ]),
    ("Data Architects", GREEN, Inches(6.9), [
        "AI/ML pipelines & feature stores",
        "Data mesh & data product mgmt",
        "Real-time streaming & EDA",
        "Data observability & DataOps",
    ]),
    ("Business Architects", PURPLE, Inches(10.1), [
        "Value stream mapping & outcomes",
        "Tech literacy (AI, cloud, platforms)",
        "OKR frameworks & CBP roadmapping",
        "Advanced facilitation techniques",
    ]),
]
for name, color, left, items in roles:
    add_shape(slide, left, Inches(1.15), Inches(2.9), Inches(0.55), fill_color=color)
    add_textbox(slide, left + Inches(0.1), Inches(1.18), Inches(2.7), Inches(0.5),
                name, font_size=15, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    txBox = add_textbox(slide, left + Inches(0.1), Inches(1.8), Inches(2.7), Inches(4.5), "", font_size=13)
    tf = txBox.text_frame
    tf.word_wrap = True
    for item in items:
        add_bullet(tf, item, font_size=13)

# =========================================================
# SLIDE 19: Cross-Role Collaboration Skills
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Cross-Role Collaboration Skills")

add_textbox(slide, Inches(1), Inches(1.15), Inches(11), Inches(0.5),
            "Required for All Architecture Roles", font_size=16, color=GRAY)

items_left = [
    "RACI framework & role interaction mapping",
    "Scaled Agile Framework (SAFe) & PI Planning",
    "JIRA: epic/story creation & linking",
    "Visual communication (Visio, Miro, Lucidchart)",
]
items_right = [
    "Architecture documentation standards & shared taxonomy",
    "Stakeholder mapping & influence strategy",
    "Design thinking & collaborative problem-solving",
]

txBox = add_textbox(slide, Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.5), "", font_size=16)
tf = txBox.text_frame
tf.word_wrap = True
for item in items_left:
    add_bullet(tf, item, font_size=16)

txBox2 = add_textbox(slide, Inches(6.8), Inches(1.8), Inches(5.8), Inches(4.5), "", font_size=16)
tf2 = txBox2.text_frame
tf2.word_wrap = True
for item in items_right:
    add_bullet(tf2, item, font_size=16)

# =========================================================
# SLIDE 20: Needed Inputs — Corporate Strategy
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Needed Inputs: From Corporate Strategy")

data = [
    ["Input", "Description"],
    ["Vision, Mission & Strategic Themes", "Clear written statements of vision, mission, strategic pillars,\nand multi-year goals traceable to capabilities & roadmaps"],
    ["Corporate & Portfolio OKRs", "Top-level and portfolio-level OKRs for alignment of tech\nroadmaps, guardrails, and enablers"],
    ["Strategic Initiatives & Bets", "Approved/emerging initiatives including target segments,\nvalue propositions, and expected outcomes"],
    ["Operating Model & Value Streams", "Current and target operating model so EA can map\ntechnology to where value flows"],
    ["Investment & Funding Guardrails", "Constraints, risk appetite, \"must win\" areas to shape\narchitectural runway"],
    ["Risk, Compliance & Policy", "Risk appetite, regulatory hot spots, security & AI ethics\nstances for EA principles"],
    ["Market & Competitive Insights", "Trends and disruption themes informing the\ntechnology radar"],
]
add_table(slide, len(data), 2, data,
          Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.8),
          col_widths=[Inches(4.0), Inches(8.3)],
          font_size=12)

# =========================================================
# SLIDE 21: Needed Inputs — Portfolio Leadership
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Needed Inputs: From Portfolio Leadership")

data = [
    ["Input", "Description"],
    ["Portfolio Backlog & Roadmap", "Prioritized portfolio backlog and 1\u20133 increment roadmap\nso EA can time enablers & platform work"],
    ["Demand Funnel & Intake", "Visibility into new requests, steering decisions, and\npipeline to decide EA engagement"],
    ["Business Cases & Constraints", "Benefits, cost, timeline, dependencies for architecture\nimpact assessments"],
    ["Solution & Platform Landscape", "Key platforms, solutions in flight, tech debt hotspots,\nand known constraints"],
    ["Risks, Issues & Dependencies", "Cross-program dependencies and capacity constraints\nfor architectural options"],
    ["PI Objectives & Outcomes", "Planned vs. actual results to refine EA roadmaps and\nadjust guardrails"],
]
add_table(slide, len(data), 2, data,
          Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.0),
          col_widths=[Inches(4.0), Inches(8.3)],
          font_size=12)

# =========================================================
# SLIDE 22: Closing
# =========================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(2.5), prs.slide_width, Inches(3.0), fill_color=RED)
add_textbox(slide, Inches(1), Inches(2.8), Inches(11), Inches(1.0),
            "Thank You", font_size=44, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1), Inches(3.9), Inches(11), Inches(0.8),
            "Red Hat Architecture \u2014 2026", font_size=28, color=WHITE, alignment=PP_ALIGN.CENTER)

# =========================================================
# SAVE
# =========================================================
output_path = r"c:\Users\jzirkle\.cursor\CursorTest\RedHat_Architecture_30Day_Scope.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
